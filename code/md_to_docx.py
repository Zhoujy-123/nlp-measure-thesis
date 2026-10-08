# -*- coding: utf-8 -*-
"""将论文 Markdown 转为排版规范的 .docx（标题/正文/表格/引用）。"""
import re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC = "/workspace/nlp-measure-thesis/paper/毕业论文_基于MD&A文本的企业技术创新导向测度构建.md"
OUT = "/workspace/nlp-measure-thesis/paper/毕业论文_基于MD&A文本的企业技术创新导向测度构建.docx"

doc = Document()
# 正文默认字体
style = doc.styles["Normal"]
style.font.name = "SimSun"
style.font.size = Pt(12)
style.element.rPr.rFonts.set(__import__("docx").oxml.ns.qn("w:eastAsia"), "SimSun")

def add_runs(par, text):
    """处理 **加粗** 与 $公式$（转为斜体普通文本）。"""
    text = text.replace("\n", " ")
    # 先拆分公式
    parts = re.split(r"(\$[^$]+\$)", text)
    for p in parts:
        if not p:
            continue
        if p.startswith("$") and p.endswith("$"):
            r = par.add_run(p.strip("$"))
            r.italic = True
            continue
        # 加粗
        segs = re.split(r"(\*\*[^*]+\*\*)", p)
        for s in segs:
            if not s:
                continue
            if s.startswith("**") and s.endswith("**"):
                r = par.add_run(s[2:-2]); r.bold = True
            else:
                r = par.add_run(s)

with open(SRC, encoding="utf-8") as f:
    lines = f.read().splitlines()

i = 0
while i < len(lines):
    line = lines[i].rstrip()
    if not line.strip():
        i += 1; continue
    # 标题
    m = re.match(r"^(#{1,4})\s+(.*)$", line)
    if m:
        level = len(m.group(1))
        h = doc.add_heading(m.group(2), level=min(level, 3))
        i += 1; continue
    # 引用块
    if line.startswith(">"):
        par = doc.add_paragraph()
        par.paragraph_format.left_indent = Pt(18)
        add_runs(par, line.lstrip("> ").strip())
        i += 1; continue
    # 表格（连续以 | 开头的行）
    if line.startswith("|"):
        tbl_rows = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            # 跳过分隔行
            if set("".join(cells)) <= set("-: "):
                i += 1; continue
            tbl_rows.append(cells)
            i += 1
        if tbl_rows:
            ncol = max(len(r) for r in tbl_rows)
            table = doc.add_table(rows=len(tbl_rows), cols=ncol)
            table.style = "Light Grid Accent 1"
            for ri, row in enumerate(tbl_rows):
                for ci in range(ncol):
                    cell = table.cell(ri, ci)
                    val = row[ci] if ci < len(row) else ""
                    val = val.replace("\\multicolumn{4}{c}{", "").replace("}", "")
                    cell.text = ""
                    add_runs(cell.paragraphs[0], val)
        continue
    # 普通段落
    par = doc.add_paragraph()
    add_runs(par, line)
    i += 1

doc.save(OUT)
print("已生成:", OUT)
print("段落数:", len(doc.paragraphs), "表格数:", len(doc.tables))
